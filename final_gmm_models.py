"""Run the locked level-rate Difference-GMM product schedule.

The estimator performs the only first difference.  This script deliberately
never passes a d_Int_Rate_* column to either backend.
"""
from __future__ import annotations

import argparse
import contextlib
from copy import copy
import hashlib
import io
import json
import math
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from scipy.stats import chi2, norm


ROOT = Path(__file__).resolve().parent
RESULTS_ROOT = ROOT / "Results" / "GMM"
SHOCK = ROOT / "mon_shock_dataset.pkl"
FEDERAL = ROOT / "Operations" / "fed_analys.xlsx"
RSCRIPT = Path(r"C:\Program Files\R\R-4.4.2\bin\x64\Rscript.exe")
R_LIBRARY = (ROOT / ".agent" / "tmp" / "r-library").resolve()
LAGS = (2, 3)
CONTROLS = ("Zakred", "CPI_reg", "Cap_to_assets")
CLUSTERS = ("Cluster_new_cd_1", "Cluster_new_cd_3", "Cluster_new_cd_4")
BASE_DIRECT = ("Mon_Shock", *(f"Mon_Shock_{name}" for name in CLUSTERS))
SCHEDULE = {
    "ConsCred": {"outcome": "Int_Rate_ConsCred", "workbook": ROOT / "Operations" / "conscred_reg_analys.xlsx", "macro": ("Conscred_MaP_Fact", "Conscred_MaP_Announcement")},
    "FL": {"outcome": "Int_Rate_FL", "workbook": ROOT / "Operations" / "fl_cred_reg_analys.xlsx", "macro": ("MaP_Fact", "MaP_Announcement")},
    "Mort": {"outcome": "Int_Rate_Mort", "workbook": ROOT / "Operations" / "mortgage_cred_reg_analys.xlsx", "macro": ("Mortgage_MaP_Fact", "Mortgage_MaP_Announcement")},
}
VARIANTS = {"Base": (), "Fact": (0,), "Announcement": (1,), "Both": (0, 1)}
FROZEN = {"rows": 6150, "groups": 75, "input_sha256": "98b0736a6860c69c96eeaa326a051df25f233dcedfdbde414cf9d1d768a418c4", "key_sha256": "b56b19be5beb66f62a66bd460fac9db181b294c8b12de5a8624581a4afdeb258", "instrument_count": 12, "instrument_rank": 12, "r_sargan_p": .3164, "hansen_p": .31643549987437725, "ar1_p": .053821204091832824, "ar2_p": .2732702102573016}

# This is the tested R pgmm contract from Modules/gmm_level_specification_search.py.
R_BACKEND = r'''args <- commandArgs(trailingOnly=TRUE)
.libPaths(c(normalizePath(args[4]), .libPaths())); library(jsonlite); library(plm)
input <- read.csv(args[1], check.names=FALSE); spec <- fromJSON(args[3]); f <- as.formula(spec$r_formula)
h <- function(x) if(inherits(x,"htest")) list(statistic=unname(x$statistic),df=unname(x$parameter),p_value=x$p.value,method=x$method) else list(error=conditionMessage(x))
out <- tryCatch({ fit <- pgmm(f,input,index=c("region_id","month_id"),effect="individual",model="twosteps",transformation="d",fsm="G",collapse=TRUE); W <- do.call(rbind,fit$W); if(ncol(W)!=length(spec$instrument_labels)) stop("R instrument labels do not match canonical specification"); colnames(W) <- unlist(spec$instrument_labels); write.csv(W,args[5],row.names=FALSE); d <- svd(W,nu=0,nv=0)$d; v<-tryCatch(vcovHC(fit),error=function(e) NULL); cv<-if(is.null(v)) NULL else list(names=colnames(v),values=unname(v)); list(status="completed",formula=spec$r_formula,nobs=nobs(fit),groups=length(fit$W),instrument_count=ncol(W),instrument_rank=qr(W)$rank,singular_values=d,instrument_labels=colnames(W),key_sha256=spec$key_sha256,input_sha256=spec$input_sha256,coefficients=as.list(coef(fit)),robust_covariance=cv,ar1=h(tryCatch(mtest(fit,order=1L,vcov=vcovHC(fit)),error=function(e)e)),ar2=h(tryCatch(mtest(fit,order=2L,vcov=vcovHC(fit)),error=function(e)e)),sargan=h(tryCatch(sargan(fit,weights="twosteps"),error=function(e)e)),sargan_label="classical_heteroskedasticity_sensitive") },error=function(e) list(status="failed",error=conditionMessage(e),formula=spec$r_formula))
write_json(out,args[2],auto_unbox=TRUE,null="null",na="null",digits=16)'''

PY_BACKEND = r'''import contextlib,io,json,sys,hashlib,inspect
import numpy as np,pandas as pd
from systemgmmkit.pydynpd_backend import _apply_numpy_compatibility_shims
from pydynpd import regression,specification_tests
data=pd.read_csv(sys.argv[1]); spec=json.load(open(sys.argv[2],encoding="utf8")); out=io.StringIO()
try:
 _apply_numpy_compatibility_shims()
 with contextlib.redirect_stdout(out): fitted=regression.abond(spec["pydynpd_command"],data,["Region","Date"])
 model=fitted.models[0]; info=model.z_information; height=int(model.z_list.shape[0]/model.N); z=np.concatenate([model.z_list[i*height:(i+1)*height,:] for i in range(model.N)],axis=1).T; np.savetxt(sys.argv[4],z,delimiter=","); s=np.linalg.svd(z,compute_uv=False); tol=float(np.finfo(float).eps*max(z.shape)*s[0]) if len(s) else 0.; ar={x.lag:x.P_value for x in model.AR_list}; h=model.hansen
 if z.shape[1]!=len(spec["instrument_labels"]): raise ValueError("Python instrument labels do not match canonical specification")
 result={"status":"completed","raw_console":out.getvalue(),"groups":int(model.N),"instrument_count":int(info.num_instr),"instrument_rank":int((s>tol).sum()),"singular_values":s.tolist(),"instrument_labels":spec["instrument_labels"],"key_sha256":spec["key_sha256"],"input_sha256":spec["input_sha256"],"ar1_p":ar.get(1),"ar2_p":ar.get(2),"hansen_statistic":getattr(h,"test_value",None),"hansen_df":getattr(h,"df",None),"hansen_p":getattr(h,"p_value",None),"hansen_robustness":"two_step_empirical_moment_covariance","hansen_source_sha256":hashlib.sha256(inspect.getsource(specification_tests.hansen_overid).encode()).hexdigest()}
except BaseException as e: result={"status":"failed","raw_console":out.getvalue(),"error_type":type(e).__name__,"error":str(e)}
json.dump(result,open(sys.argv[3],"w",encoding="utf8"),ensure_ascii=False)'''


def digest(value: Path | pd.DataFrame) -> str:
    raw = value.read_bytes() if isinstance(value, Path) else value.to_csv(index=False, date_format="%Y-%m-%d").encode()
    return hashlib.sha256(raw).hexdigest()


def formula(outcome: str, direct: tuple[str, ...]) -> tuple[str, str, list[str]]:
    labels = [f"gmm:{outcome}:L{lag}" for lag in LAGS] + [f"gmm:{name}:L{lag}" for name in CONTROLS for lag in LAGS] + [f"iv:{name}" for name in direct]
    r = f"{outcome} ~ lag({outcome}, 1) + {' + '.join((*CONTROLS, *direct))} | lag({outcome}, 2:3) + " + " + ".join(f"lag({name}, 2:3)" for name in CONTROLS) + f" | {' + '.join(direct)}"
    py = f"{outcome} L1.{outcome} {' '.join((*CONTROLS, *direct))} | gmm({outcome} {' '.join(CONTROLS)}, 2:3) iv({' '.join(direct)}) | collapse nolevel"
    return r, py, labels


def monthly_panel_check(data: pd.DataFrame) -> None:
    if data.duplicated(["Region", "Date"]).any():
        raise ValueError("duplicate Region-Date keys")
    expected = pd.date_range(data.Date.min(), data.Date.max(), freq="MS")
    if any(not pd.DatetimeIndex(x).equals(expected) for _, x in data.sort_values("Date").groupby("Region")["Date"]):
        raise ValueError("panel is not balanced consecutive monthly data")
    if data.groupby("Region").size().nunique() != 1:
        raise ValueError("panel is unbalanced")


def load_product(product: str) -> tuple[pd.DataFrame | None, dict[str, Any], str | None]:
    info = SCHEDULE[product]; outcome, workbook = info["outcome"], info["workbook"]
    data = pd.read_excel(workbook).drop(columns=lambda x: str(x).startswith("Unnamed:"), errors="ignore")
    if outcome not in data:
        return None, {"product": product, "workbook": str(workbook), "workbook_sha256": digest(workbook), "outcome": outcome}, "blocked_missing_level_outcome: source lacks " + outcome
    required = {"Region", "Date", outcome, *CONTROLS, *CLUSTERS}
    missing = sorted(required.difference(data.columns))
    if missing:
        raise ValueError(f"{product}: missing regional columns {missing}")
    data["Date"] = pd.to_datetime(data["Date"], errors="raise").dt.to_period("M").dt.to_timestamp()
    monthly_panel_check(data)
    clusters = data[list(CLUSTERS)]
    if not clusters.isin((0, 1)).all().all() or not clusters.sum(axis=1).isin((0, 1)).all():
        raise ValueError(f"{product}: cluster dummies must be mutually exclusive 0/1")
    shock = pd.read_pickle(SHOCK).copy(); shock["Date"] = pd.to_datetime(shock["Date"], errors="raise").dt.to_period("M").dt.to_timestamp()
    if shock.duplicated("Date").any() or not {"Date", "Mon_Shock"}.issubset(shock):
        raise ValueError("shock source requires unique Date/Mon_Shock")
    macro = pd.read_excel(FEDERAL, usecols=lambda x: x == "Date" or x in info["macro"]).drop(columns=lambda x: str(x).startswith("Unnamed:"), errors="ignore")
    macro["Date"] = pd.to_datetime(macro["Date"], errors="raise").dt.to_period("M").dt.to_timestamp()
    if macro.duplicated("Date").any() or set(info["macro"]).difference(macro):
        raise ValueError(f"{product}: federal Macropru source is incomplete or duplicated")
    data = data.merge(shock[["Date", "Mon_Shock"]], on="Date", how="left", validate="many_to_one").merge(macro, on="Date", how="left", validate="many_to_one")
    if data[["Mon_Shock", *info["macro"]]].isna().any().any():
        raise ValueError(f"{product}: source merge misses shock or Macropru months")
    for name in CLUSTERS:
        data[f"Mon_Shock_{name}"] = data["Mon_Shock"] * data[name]
    data = data.sort_values(["Region", "Date"]).reset_index(drop=True)
    required_values = [outcome, *CONTROLS, *BASE_DIRECT, *info["macro"]]
    if data[required_values].isna().any().any():
        raise ValueError(f"{product}: required model values are missing; no complete-case row removal is allowed")
    audit = {"product": product, "outcome": outcome, "workbook": str(workbook), "workbook_sha256": digest(workbook), "federal_workbook": str(FEDERAL), "federal_sha256": digest(FEDERAL), "shock": str(SHOCK), "shock_sha256": digest(SHOCK), "merged_sha256": digest(data), "rows": len(data), "regions": int(data.Region.nunique()), "dates": int(data.Date.nunique()), "window": f"{data.Date.min():%Y-%m} to {data.Date.max():%Y-%m}", "missing_values": int(data[[outcome, *CONTROLS, *BASE_DIRECT, *info['macro']]].isna().sum().sum())}
    return data, audit, None


def approved_python() -> Path:
    probe = "from importlib.metadata import version; print(version('pydynpd'))"
    for candidate in sorted((ROOT / ".uv-cache" / "archive-v0").glob("*/Scripts/python.exe")):
        result = subprocess.run([str(candidate), "-c", probe], capture_output=True, text=True, check=False)
        if result.returncode == 0 and result.stdout.strip() == "0.2.2":
            return candidate
    raise RuntimeError("approved pydynpd 0.2.2 interpreter is unavailable")


def run_r(spec: dict[str, Any], run_dir: Path) -> dict[str, Any]:
    stem = spec["id"]; inp, result, conf, matrix = (run_dir / f"{stem}_r_input.csv", run_dir / f"{stem}_r.json", run_dir / f"{stem}_spec.json", run_dir / f"{stem}_r_Z.csv")
    data = spec["input"].copy(); data["region_id"] = pd.factorize(data.Region)[0] + 1; data["month_id"] = pd.factorize(data.Date)[0] + 1
    data.to_csv(inp, index=False); conf.write_text(json.dumps({k: v for k, v in spec.items() if k != "input"}, default=str), encoding="utf8")
    call = subprocess.run([str(RSCRIPT), "--vanilla", str(run_dir / "r_backend.R"), str(inp), str(result), str(conf), str(R_LIBRARY), str(matrix)], capture_output=True, text=True, check=False)
    out = json.loads(result.read_text(encoding="utf8")) if result.exists() else {"status": "failed", "error": "R backend produced no JSON"}
    out["raw_console"] = call.stdout + call.stderr; out["matrix_path"] = str(matrix) if matrix.exists() else None
    return out


def run_python(spec: dict[str, Any], run_dir: Path, python: Path) -> dict[str, Any]:
    stem = spec["id"]; inp, conf, result, matrix = (run_dir / f"{stem}_py_input.csv", run_dir / f"{stem}_py_spec.json", run_dir / f"{stem}_py.json", run_dir / f"{stem}_py_Z.csv")
    spec["input"].to_csv(inp, index=False); conf.write_text(json.dumps({k: v for k, v in spec.items() if k != "input"}, default=str), encoding="utf8")
    call = subprocess.run([str(python), str(run_dir / "pydynpd_backend.py"), str(inp), str(conf), str(result), str(matrix)], capture_output=True, text=True, check=False)
    out = json.loads(result.read_text(encoding="utf8")) if result.exists() else {"status": "failed", "error": "Python backend produced no JSON"}
    out["raw_console"] = out.get("raw_console", "") + call.stdout + call.stderr; out["matrix_path"] = str(matrix) if matrix.exists() else None
    return out


def parity(r: dict[str, Any], py: dict[str, Any], spec: dict[str, Any]) -> dict[str, Any]:
    checks = {"formula": r.get("formula") == spec["r_formula"], "groups": r.get("groups") == py.get("groups") == spec["groups"], "instrument_count": r.get("instrument_count") == py.get("instrument_count") == len(spec["instrument_labels"]), "instrument_rank": r.get("instrument_rank") == py.get("instrument_rank"), "hashes": r.get("key_sha256") == py.get("key_sha256") == spec["key_sha256"] and r.get("input_sha256") == py.get("input_sha256") == spec["input_sha256"], "labels": r.get("instrument_labels") == py.get("instrument_labels") == spec["instrument_labels"]}
    try:
        rz = pd.read_csv(r["matrix_path"]).to_numpy(float); pz = np.loadtxt(py["matrix_path"], delimiter=",")
        checks["matrix_shape"] = rz.shape == pz.shape
        order = [next((j for j in range(pz.shape[1]) if np.allclose(rz[:, i], pz[:, j], rtol=1e-8, atol=1e-10)), -1) for i in range(rz.shape[1])] if checks["matrix_shape"] else []
        checks["matrix_values"] = bool(order and -1 not in order and len(set(order)) == len(order) and np.allclose(rz, pz[:, order], rtol=1e-8, atol=1e-10))
    except (KeyError, OSError, ValueError):
        checks["matrix_values"] = False
    return {"status": "match" if all(checks.values()) else "mismatch", "checks": checks}


def coefficient_rows(item: dict[str, Any]) -> list[dict[str, Any]]:
    r, cov = item.get("r", {}), item.get("r", {}).get("robust_covariance") or {}
    names, values, coefs = cov.get("names") or [], cov.get("values"), r.get("coefficients") or {}
    if values is None: return []
    matrix = np.asarray(values, dtype=float); out = []
    for i, name in enumerate(names):
        estimate, se = float(coefs.get(name, np.nan)), math.sqrt(max(float(matrix[i, i]), 0.0))
        z = estimate / se if se else np.nan; p = 2 * norm.sf(abs(z)) if np.isfinite(z) else np.nan
        out.append({"model": item["model"], "product": item["product"], "term": name, "estimate": estimate, "robust_se": se, "z_statistic": z, "p_value": p, "ci_95_low": estimate - 1.96 * se, "ci_95_high": estimate + 1.96 * se, "significance": "***" if p < .01 else "**" if p < .05 else "*" if p < .1 else ""})
    return out


def wald_rows(item: dict[str, Any]) -> list[dict[str, Any]]:
    cov, coefs = item.get("r", {}).get("robust_covariance") or {}, item.get("r", {}).get("coefficients") or {}
    names, values = cov.get("names") or [], cov.get("values")
    if values is None: return []
    matrix = np.asarray(values, dtype=float); index = {name: i for i, name in enumerate(names)}
    lag = next((x for x in names if "lag(" in x), None); direct = item["direct"]
    blocks = {"All non-lagged regressors": [x for x in names if x != lag], "Shock and cluster interactions": list(BASE_DIRECT), "Regional controls": list(CONTROLS), "Macropru block": list(item["macro_terms"])}
    rows = []
    for label, terms in blocks.items():
        terms = [x for x in terms if x in index]
        if not terms: continue
        pos = [index[x] for x in terms]; b = np.array([coefs[x] for x in terms], dtype=float); v = matrix[np.ix_(pos, pos)]
        statistic = float(b @ np.linalg.pinv(v) @ b); p = float(chi2.sf(statistic, len(terms)))
        rows.append({"model": item["model"], "product": item["product"], "block": label, "terms": ", ".join(terms), "wald_chi_square": statistic, "df": len(terms), "p_value": p, "target": "robust Wald chi-square; report, not a validity gate"})
    return rows


def regressor_diagnostics(spec: dict[str, Any]) -> list[dict[str, Any]]:
    outcome = spec["outcome"]
    data = spec["input"].sort_values(["Region", "Date"])
    rhs = [f"L1.{outcome}", *CONTROLS, *spec["direct"]]
    d = data.groupby("Region")[[*CONTROLS, *spec["direct"]]].diff()
    d.insert(0, f"L1.{outcome}", data.groupby("Region")[outcome].shift(1).groupby(data["Region"]).diff())
    d = d.dropna()
    x = d.to_numpy(float); rank = int(np.linalg.matrix_rank(x)); condition = float(np.linalg.cond(x)) if len(x) else np.nan
    rows = []
    for i, name in enumerate(rhs):
        others = np.delete(x, i, axis=1); target = x[:, i]
        fitted = others @ np.linalg.lstsq(others, target, rcond=None)[0] if others.size else np.zeros_like(target)
        denom = float(((target - target.mean()) ** 2).sum()); r2 = 1 - float(((target - fitted) ** 2).sum()) / denom if denom else 1.0
        vif = np.inf if r2 >= 1 - 1e-12 else 1 / (1 - r2)
        rows.append({"model": spec["id"], "term": name, "transformed_regressor_rank": rank, "regressor_count": len(rhs), "condition_number": condition, "vif": vif, "target": "full rank hard; VIF <= 10 and condition <= 30 descriptive", "warning": "rank failure" if rank < len(rhs) else "VIF > 10" if vif > 10 else "condition > 30" if condition > 30 else ""})
    return rows


def gate(item: dict[str, Any]) -> tuple[str, list[str], list[str]]:
    r, py, check = item.get("r", {}), item.get("py", {}), item.get("parity", {})
    failures, warnings = [], []
    if r.get("status") != "completed" or py.get("status") != "completed": failures.append("R or Python backend did not complete")
    elif check.get("status") != "match": failures.append("R/Python realised matrix parity failed")
    if r.get("instrument_rank") != r.get("instrument_count") or (r.get("instrument_count") or 0) >= item["groups"]: failures.append("instrument rank/count gate failed (full rank and K < N regions required)")
    if r.get("sargan", {}).get("p_value") is None or r["sargan"]["p_value"] < .05: failures.append("R Sargan p < .05")
    if py.get("hansen_p") is None or py["hansen_p"] < .05: failures.append("Python Hansen p < .05")
    if py.get("ar2_p") is None or py["ar2_p"] < .05: failures.append("Python AR(2) p < .05")
    diag = item.get("regressor_diagnostics", [])
    if diag and diag[0]["transformed_regressor_rank"] < diag[0]["regressor_count"]: failures.append("transformed regressor matrix is not full rank")
    if (r.get("instrument_count") or 0) / max(item["groups"], 1) >= .5: warnings.append("K/N regions >= .50")
    if py.get("ar1_p") is not None and py["ar1_p"] >= .05: warnings.append("AR(1) p >= .05; differenced residual autocorrelation was not detected")
    if any((value or 0) > .90 for value in (r.get("sargan", {}).get("p_value"), py.get("hansen_p"))): warnings.append("overidentification p > .90")
    if any(x["warning"] in {"VIF > 10", "condition > 30"} for x in diag): warnings.append("descriptive multicollinearity warning")
    return ("valid" if not failures else "invalid"), failures, warnings


def baseline_check(item: dict[str, Any]) -> list[str]:
    if item["product"] != "ConsCred" or item["variant"] != "Base": return []
    r, py = item["r"], item["py"]; actual = {"rows": item["rows"], "groups": item["groups"], "input_sha256": item["input_sha256"], "key_sha256": item["key_sha256"], "instrument_count": r.get("instrument_count"), "instrument_rank": r.get("instrument_rank"), "r_sargan_p": r.get("sargan", {}).get("p_value"), "hansen_p": py.get("hansen_p"), "ar1_p": py.get("ar1_p"), "ar2_p": py.get("ar2_p")}
    failed = []
    for key, expected in FROZEN.items():
        value = actual.get(key); tolerance = .00005 if key == "r_sargan_p" else 1e-6 if key in {"hansen_p", "ar1_p", "ar2_p"} else 0
        if (not isinstance(expected, float) and value != expected) or (isinstance(expected, float) and (value is None or abs(value - expected) >= tolerance)):
            failed.append(f"{key}: expected {expected}, got {value}")
    return failed


def fit(product: str, variant: str, data: pd.DataFrame, audit: dict[str, Any], run_dir: Path, python: Path) -> dict[str, Any]:
    outcome, macro = SCHEDULE[product]["outcome"], SCHEDULE[product]["macro"]
    macro_terms = tuple(macro[i] for i in VARIANTS[variant]); direct = (*BASE_DIRECT, *macro_terms)
    if any(data.groupby("Date")[name].first().diff().dropna().eq(0).all() for name in macro_terms):
        return blocked_item(product, variant, outcome, audit, "blocked_macropru_no_transformed_variation")
    # Keep the audited column order used by the passing consumer-credit run:
    # outcome, direct terms, then endogenous controls.
    columns = ["Region", "Date", outcome, *direct, *CONTROLS]
    sample = data[columns].sort_values(["Region", "Date"]).reset_index(drop=True)
    monthly_panel_check(sample)
    r_formula, py_command, labels = formula(outcome, direct); model = f"{product}_{variant}"
    spec = {"id": model, "outcome": outcome, "input": sample, "groups": int(sample.Region.nunique()), "rows": len(sample), "key_sha256": digest(sample[["Region", "Date"]]), "input_sha256": digest(sample), "r_formula": r_formula, "pydynpd_command": py_command, "instrument_labels": labels, "direct": direct}
    r, py = run_r(spec, run_dir), run_python(spec, run_dir, python); item = {**spec, "product": product, "variant": variant, "model": model, "macro_terms": macro_terms, "r": r, "py": py, "audit": audit}
    item["parity"] = parity(r, py, spec); item["regressor_diagnostics"] = regressor_diagnostics(spec); item["status"], item["failures"], item["warnings"] = gate(item)
    reproduce = baseline_check(item)
    if reproduce: item["status"] = "baseline_reproduction_failure"; item["failures"].append("baseline reproduction failure: " + "; ".join(reproduce))
    return item


def blocked_item(product: str, variant: str, outcome: str, audit: dict[str, Any], reason: str) -> dict[str, Any]:
    return {"product": product, "variant": variant, "model": f"{product}_{variant}", "outcome": outcome, "status": reason.split(":")[0], "failures": [reason], "warnings": [], "audit": audit, "rows": None, "groups": None, "macro_terms": (), "direct": (), "r": {}, "py": {}, "parity": {}, "regressor_diagnostics": []}


def validation_rows(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for x in items:
        r, py = x["r"], x["py"]
        rows.append({"model": x["model"], "product": x["product"], "variant": x["variant"], "outcome": x["outcome"], "verdict": x["status"], "failure_or_block_reason": "; ".join(x["failures"]), "warnings": "; ".join(x["warnings"]), "rows": x["rows"], "regions": x["groups"], "formula": x.get("r_formula"), "matrix_parity": x["parity"].get("status"), "instruments_K": r.get("instrument_count"), "instrument_rank": r.get("instrument_rank"), "K_over_N_target": "< 0.50 warning at/above", "AR1_p": py.get("ar1_p"), "AR1_target": "< .05 expected warning otherwise", "AR2_p": py.get("ar2_p"), "AR2_target": ">= .05 hard", "R_Sargan_p": r.get("sargan", {}).get("p_value"), "R_Sargan_target": ">= .05 hard; heteroskedasticity-sensitive", "Python_Hansen_p": py.get("hansen_p"), "Python_Hansen_target": ">= .05 hard"})
    return rows


def comparison_rows(items: list[dict[str, Any]], data: dict[str, pd.DataFrame | None]) -> list[dict[str, Any]]:
    out = []
    for product in SCHEDULE:
        base = next(x for x in items if x["product"] == product and x["variant"] == "Base")
        for variant in ("Fact", "Announcement", "Both"):
            alt = next(x for x in items if x["product"] == product and x["variant"] == variant)
            if base["status"].startswith("blocked") or alt["status"].startswith("blocked"):
                out.append({"product": product, "comparison": f"Base vs {variant}", "status": "blocked_missing_level_outcome", "reason": "; ".join(base["failures"] + alt["failures"]), "common_sample_keys": False})
                continue
            equal = base.get("key_sha256") == alt.get("key_sha256")
            changed = "; ".join(sorted(set(alt["failures"]) - set(base["failures"]))) or "none"
            base_coef, alt_coef = base["r"].get("coefficients", {}), alt["r"].get("coefficients", {})
            coefficient_changes = {term: float(alt_coef[term]) - float(base_coef[term]) for term in sorted(set(base_coef).intersection(alt_coef))}
            out.append({"product": product, "comparison": f"Base vs {variant}", "status": "compared", "reason": changed, "common_sample_keys": equal, "changed_coefficients_variant_minus_base": json.dumps(coefficient_changes), "base_verdict": base["status"], "variant_verdict": alt["status"], "base_hansen_p": base["py"].get("hansen_p"), "variant_hansen_p": alt["py"].get("hansen_p"), "base_ar2_p": base["py"].get("ar2_p"), "variant_ar2_p": alt["py"].get("ar2_p"), "base_sargan_p": base["r"].get("sargan", {}).get("p_value"), "variant_sargan_p": alt["r"].get("sargan", {}).get("p_value")})
    return out


def write_workbook(run_dir: Path, items: list[dict[str, Any]], audits: list[dict[str, Any]], comparisons: list[dict[str, Any]]) -> None:
    coefficients = [row for item in items for row in coefficient_rows(item)]
    wald = [row for item in items for row in wald_rows(item)]
    serial, overid, instruments, diagnostics, specs = [], [], [], [], []
    for x in items:
        r, py = x["r"], x["py"]
        specs.append({"model": x["model"], "outcome": x["outcome"], "R_formula": x.get("r_formula"), "Python_command": x.get("pydynpd_command"), "endogenous_GMM_lags": f"{x.get('outcome')} and {', '.join(CONTROLS)} at 2:3, collapsed", "ordinary_IV_terms": ", ".join(x["direct"]), "Macropru_terms": ", ".join(x["macro_terms"]), "estimator": "two-step Difference GMM, individual effects, fsm=G, nolevel"})
        for order in (1, 2):
            target = "AR(1) p < .05 expected" if order == 1 else "AR(2) p >= .05 hard"
            for backend, statistic, p_value in (("R plm", r.get(f"ar{order}", {}).get("statistic"), r.get(f"ar{order}", {}).get("p_value")), ("Python pydynpd", None, py.get(f"ar{order}_p"))):
                verdict = "blocked" if p_value is None else "warning" if order == 1 and p_value >= .05 else "pass" if order == 1 or p_value >= .05 else "fail"
                serial.append({"model": x["model"], "backend": backend, "test": f"AR({order})", "statistic": statistic, "p_value": p_value, "target": target, "verdict": verdict})
        for test, statistic, df, p_value, method in (("R Sargan", r.get("sargan", {}).get("statistic"), r.get("sargan", {}).get("df"), r.get("sargan", {}).get("p_value"), "classical; heteroskedasticity-sensitive"), ("Python Hansen J", py.get("hansen_statistic"), py.get("hansen_df"), py.get("hansen_p"), "two-step empirical moment covariance")):
            overid.append({"model": x["model"], "test": test, "statistic": statistic, "df": df, "p_value": p_value, "method": method, "target": "p >= .05 hard", "verdict": "blocked" if p_value is None else "pass" if p_value >= .05 else "fail"})
        instruments.append({"model": x["model"], "instrument_count_K": r.get("instrument_count"), "instrument_rank": r.get("instrument_rank"), "regions_N": x["groups"], "K_over_N": (r.get("instrument_count") or 0) / x["groups"] if x["groups"] else None, "labels": ", ".join(r.get("instrument_labels") or []), "realised_matrix_parity": x["parity"].get("status"), "target": "full rank; K < N hard; K/N < .50 preferred"})
        diagnostics.extend(x["regressor_diagnostics"])
    sheets = {"Validation_Summary": pd.DataFrame(validation_rows(items)), "Model_Specification": pd.DataFrame(specs), "Coefficients": pd.DataFrame(coefficients), "Wald_Tests": pd.DataFrame(wald), "AB_Serial_Correlation": pd.DataFrame(serial), "Overidentification": pd.DataFrame(overid), "Instruments": pd.DataFrame(instruments), "Regressor_Diagnostics": pd.DataFrame(diagnostics), "Model_Comparison": pd.DataFrame(comparisons), "Data_Audit": pd.DataFrame(audits)}
    with pd.ExcelWriter(run_dir / "final_gmm_models.xlsx", engine="openpyxl") as writer:
        for name, table in sheets.items(): table.to_excel(writer, sheet_name=name, index=False)
    from openpyxl import load_workbook
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import PatternFill
    book = load_workbook(run_dir / "final_gmm_models.xlsx")
    for ws in book.worksheets:
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            letter = column[0].column_letter; ws.column_dimensions[letter].width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            header = str(column[0].value or "")
            number_format = "0.000000E+00" if header.endswith("_p") or header == "p_value" else "0.000000"
            if header in {"df", "rows", "regions", "instrument_rank", "instruments_K", "instrument_count_K", "regressor_count", "transformed_regressor_rank"}: number_format = "0"
            for cell in column[1:]:
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool): cell.number_format = number_format
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            font = copy(cell.font); font.color = "FFFFFF"; font.bold = True; cell.font = font
    summary = book["Validation_Summary"]; headers = {cell.value: cell.column_letter for cell in summary[1]}
    if "verdict" in headers:
        status_range = f"{headers['verdict']}2:{headers['verdict']}{summary.max_row}"
        summary.conditional_formatting.add(status_range, CellIsRule(operator="equal", formula=['"valid"'], fill=PatternFill("solid", fgColor="C6EFCE")))
        summary.conditional_formatting.add(status_range, CellIsRule(operator="equal", formula=['"invalid"'], fill=PatternFill("solid", fgColor="FFC7CE")))
        summary.conditional_formatting.add(status_range, FormulaRule(formula=[f'LEFT({headers["verdict"]}2,7)="blocked"'], fill=PatternFill("solid", fgColor="FFEB9C")))
        summary.conditional_formatting.add(status_range, CellIsRule(operator="equal", formula=['"baseline_reproduction_failure"'], fill=PatternFill("solid", fgColor="FFC7CE")))
    book.save(run_dir / "final_gmm_models.xlsx")


def self_check() -> int:
    assert list(SCHEDULE) == ["ConsCred", "FL", "Mort"] and set(VARIANTS) == {"Base", "Fact", "Announcement", "Both"}
    r, py, labels = formula("Int_Rate_ConsCred", BASE_DIRECT)
    assert "lag(Int_Rate_ConsCred, 2:3)" in r and "gmm(Int_Rate_ConsCred Zakred CPI_reg Cap_to_assets, 2:3)" in py and len(labels) == 12
    assert "d_Int_Rate" not in r + py
    mort = pd.read_excel(SCHEDULE["Mort"]["workbook"], nrows=1)
    assert "Int_Rate_Mort" in mort and "d_Int_Rate_Mort" in mort
    panel, _, reason = load_product("Mort")
    assert panel is not None and reason is None
    sample = {"product": "x", "variant": "Base", "r": {"status": "completed", "instrument_count": 12, "instrument_rank": 12, "sargan": {"p_value": .2}}, "py": {"status": "completed", "hansen_p": .2, "ar1_p": .2, "ar2_p": .2}, "parity": {"status": "match"}, "groups": 75, "regressor_diagnostics": [{"transformed_regressor_rank": 2, "regressor_count": 2, "warning": ""}]}
    assert gate(sample)[0] == "valid" and "AR(1)" in gate(sample)[2][0]
    print("self-check passed")
    return 0


def main(argv: list[str] | None = None) -> int:
    args = argparse.ArgumentParser(); args.add_argument("--self-check", action="store_true"); parsed = args.parse_args(argv)
    if parsed.self_check: return self_check()
    run_dir = RESULTS_ROOT / f"final_product_level_models_{datetime.now():%Y%m%d_%H%M%S}"; run_dir.mkdir(parents=True, exist_ok=False)
    (run_dir / "r_backend.R").write_text(R_BACKEND, encoding="utf8"); (run_dir / "pydynpd_backend.py").write_text(PY_BACKEND, encoding="utf8")
    try: python = approved_python()
    except RuntimeError as error: print(error, file=sys.stderr); return 2
    items, audits, panels = [], [], {}
    for product, info in SCHEDULE.items():
        try: panel, audit, block = load_product(product)
        except (OSError, ValueError) as error: panel, audit, block = None, {"product": product, "outcome": info["outcome"]}, "blocked_data_contract: " + str(error)
        panels[product] = panel; audits.append({**audit, "block_reason": block or ""})
        for variant in VARIANTS:
            items.append(blocked_item(product, variant, info["outcome"], audit, block) if block else fit(product, variant, panel, audit, run_dir, python))
    for item in items:
        for backend in ("r", "py"):
            if item.get(backend): (run_dir / f"{item['model']}_{backend}_console.log").write_text(item[backend].get("raw_console", ""), encoding="utf8")
    comparisons = comparison_rows(items, panels); write_workbook(run_dir, items, audits, comparisons)
    (run_dir / "diagnostics.json").write_text(json.dumps({"created": datetime.now().isoformat(), "items": items, "comparisons": comparisons}, ensure_ascii=False, indent=2, default=str), encoding="utf8")
    (run_dir / "report.md").write_text("# Final product-level Difference-GMM run\n\n| Model | Verdict | Reason |\n|---|---|---|\n" + "\n".join(f"| {x['model']} | {x['status']} | {'; '.join(x['failures']) or '—'} |" for x in items) + "\n", encoding="utf8")
    print(f"bundle={run_dir}")
    return 0


if __name__ == "__main__": raise SystemExit(main())
